import zipfile
from io import BytesIO

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from bidreader.app import _validate_evidence, app, settings
from bidreader.database import Base, get_session
from bidreader.evaluation import aggregate_corpus, evaluate
from bidreader.extraction import _package_specific_fields, extract_criteria, extract_fields
from bidreader.models import Analysis, Run, RunEvent
from bidreader.parsers import _ocr_page, parse_document, parse_docx


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("第一章 项目概况", level=1)
    document.add_paragraph("项目名称：城市供水设施改造工程")
    document.add_paragraph("招标编号：ZB-2026-015")
    document.add_heading("第二章 评标办法", level=1)
    document.add_paragraph("技术评分标准")
    document.add_paragraph("类似工程业绩：每提供一项得 2 分，最高 6 分。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目负责人"
    table.cell(0, 1).text = "具备高级工程师职称得 4 分"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _blocks():
    return [
        {"block_id": "b-0", "kind": "heading", "text": "评标办法", "page_no": 3, "bbox": None, "section_path": ["评标办法"], "source_index": 0},
        {"block_id": "b-1", "kind": "line", "text": "评标办法前附表之四：技术评分标准", "page_no": 3, "bbox": None, "section_path": ["评标办法"], "source_index": 1},
        {"block_id": "b-head", "kind": "table_row", "text": "序号 | 项目 | 评审内容及分值 | 项目内容", "cells": ["序号", "项目", "评审内容及分值", "项目内容"], "page_no": 4, "bbox": [30, 20, 400, 35], "section_path": ["评标办法"], "source_index": 2, "table_id": "t-1", "row_index": 0},
        {"block_id": "b-2", "kind": "table_row", "text": "1 | 相似工程业绩 | 相似工程业绩（最高 6 分） | 每提供一项得 2 分，最高 6 分", "cells": ["1", "相似工程业绩", "相似工程业绩（最高 6 分）", "每提供一项得 2 分，最高 6 分"], "page_no": 4, "bbox": [30, 40, 400, 80], "section_path": ["评标办法", "技术评分标准"], "source_index": 3, "table_id": "t-1", "row_index": 1},
        {"block_id": "b-3", "kind": "heading", "text": "投标文件格式", "page_no": 15, "bbox": None, "section_path": ["投标文件格式"], "source_index": 4},
        {"block_id": "b-4", "kind": "line", "text": "未响应投标文件格式的，不予受理", "page_no": 16, "bbox": None, "section_path": ["投标文件格式"], "source_index": 5},
    ]


def test_docx_preserves_body_order_headings_and_table_rows(tmp_path):
    path = tmp_path / "sample.docx"
    path.write_bytes(_docx_bytes())

    blocks, sections = parse_docx(path)

    texts = [block.text for block in blocks]
    assert texts.index("项目名称：城市供水设施改造工程") < texts.index("技术评分标准")
    assert any(block.kind == "table_row" and "项目负责人" in block.text for block in blocks)
    assert any(section["title"] == "第一章 项目概况" for section in sections)


def test_project_fields_are_grounded_and_not_invented():
    fields = extract_fields(_blocks(), "doc-1")
    by_name = {field.field: field for field in fields}

    # The test blocks deliberately contain no project identifier or schedule.
    assert by_name["project_title"].state == "not_found"
    assert by_name["project_number"].state == "not_found"
    assert by_name["price_limit"].normalized_value is None


def test_project_field_candidate_carries_exact_source_evidence():
    blocks = [{"block_id": "b-title", "kind": "line", "text": "项目名称：城市供水设施改造工程",
              "page_no": 2, "bbox": [30, 40, 300, 55], "section_path": ["项目概况"], "source_index": 0}]

    fields = {item.field: item for item in extract_fields(blocks, "doc-2")}

    assert fields["project_title"].raw_value == "城市供水设施改造工程"
    assert fields["project_title"].evidence[0].quote == blocks[0]["text"]
    assert fields["project_title"].evidence[0].page_no == 2
    assert "Inupedia/tender-extract" in fields["project_title"].review_reason


def test_vendored_tender_extract_patterns_add_a_project_number_candidate():
    blocks = [
        {"block_id": "b-name", "kind": "paragraph", "text": "项目名称：城市供水设施改造工程",
         "page_no": None, "bbox": None, "section_path": [], "source_index": 0},
        {"block_id": "b-id", "kind": "paragraph", "text": "招标编号：ZB-2026-015",
         "page_no": None, "bbox": None, "section_path": [], "source_index": 1},
    ]

    fields = {item.field: item for item in extract_fields(blocks, "doc-oss")}

    assert fields["project_title"].raw_value == "城市供水设施改造工程"
    assert fields["project_number"].raw_value == "ZB-2026-015"
    assert fields["project_number"].evidence[0].block_id == "b-id"


def test_vendored_amount_normalization_is_only_mapped_with_explicit_tender_label():
    blocks = [
        {"block_id": "b-price", "kind": "paragraph", "text": "最高投标限价：500万元",
         "page_no": 2, "bbox": [10, 20, 160, 40], "section_path": [], "source_index": 0},
        {"block_id": "b-unlabeled", "kind": "paragraph", "text": "本项目服务费用约500万元",
         "page_no": 3, "bbox": [10, 20, 160, 40], "section_path": [], "source_index": 1},
    ]

    fields = {item.field: item for item in extract_fields(blocks, "doc-price")}

    assert fields["price_limit"].raw_value == "500万元"
    assert fields["price_limit"].normalized_value == "5000000.00"
    assert fields["price_limit"].evidence[0].block_id == "b-price"


def test_field_extraction_rejects_cross_references_and_non_date_deadline_text():
    blocks = [
        {"block_id": "b-title", "kind": "paragraph", "text": "项目名称：见投标人须知前附表。",
         "page_no": None, "bbox": None, "section_path": [], "source_index": 0},
        {"block_id": "b-date", "kind": "paragraph", "text": "投标截止之日至中标通知书送达前，均适用该规定。",
         "page_no": None, "bbox": None, "section_path": [], "source_index": 1},
        {"block_id": "b-duration", "kind": "paragraph", "text": "计划工期要求以技术规范书和工程量清单为准。",
         "page_no": None, "bbox": None, "section_path": [], "source_index": 2},
    ]

    fields = {item.field: item for item in extract_fields(blocks, "doc-field-quality")}

    assert fields["project_title"].state == "not_found"
    assert fields["submission_deadline"].state == "not_found"
    assert fields["duration"].state == "not_found"


def test_evidence_validation_rejects_forged_page_or_block_quote():
    document = {"document_id": "doc-1", "blocks": [
        {"block_id": "b-1", "text": "项目名称：样例工程", "page_no": 3, "bbox": [10, 20, 90, 40]}]}
    evidence = [{"document_id": "doc-1", "page_no": 3, "block_id": "b-1", "quote": "样例工程",
                 "bbox": [10, 20, 90, 40]}]

    assert _validate_evidence(document, evidence, "样例工程")
    assert not _validate_evidence(document, [{**evidence[0], "page_no": 4}], "样例工程")
    assert not _validate_evidence(document, [{**evidence[0], "quote": "虚构工程"}], "虚构工程")


def test_evidence_validation_checks_xlsx_sheet_and_cell_anchor():
    document = {"document_id": "doc-xlsx", "blocks": [{
        "block_id": "b-row", "text": "包1 | 保证金581万元", "page_no": None, "bbox": None,
        "sheet_name": "保证金清单", "cell_range": "A2:B2"}]}
    evidence = [{"document_id": "doc-xlsx", "page_no": None, "bbox": None, "block_id": "b-row",
                 "sheet_name": "保证金清单", "cell_range": "A2:B2", "quote": "保证金581万元"}]

    assert _validate_evidence(document, evidence, "保证金581万元")
    assert not _validate_evidence(document, [{**evidence[0], "cell_range": "A3:B3"}], "保证金581万元")


def test_paddle_ocr_boxes_map_back_to_pdf_points(monkeypatch):
    class FakeResult:
        def __init__(self):
            self.json = {"res": {"rec_texts": ["最高限价"], "rec_boxes": [[20, 40, 180, 80]], "rec_scores": [0.98]}}

    class FakePipeline:
        def predict(self, _image):
            return [FakeResult()]

    monkeypatch.setattr("bidreader.parsers._paddle_ocr", lambda: FakePipeline())
    blocks = _ocr_page(Image.new("RGB", (360, 720)), 2, 180, 360, 7)

    assert blocks[0].text == "最高限价"
    assert blocks[0].bbox == [10, 20, 90, 40]
    assert blocks[0].page_no == 2


def test_gold_evaluation_reports_recall_precision_score_and_citation_accuracy():
    digest = "a" * 64
    gold = {
        "schema_version": "1.0",
        "annotation_complete": True,
        "document": {"sha256": digest},
        "coverage_reviewed": ["project_fields", "technical_criteria", "commercial_criteria", "price_criteria"],
        "project_fields": {"project_title": {"value": "示例 工程"}, "project_number": {"value": None}},
        "criteria": [{"category": "technical", "source_quote": "类似工程业绩每项得2分最高6分。",
                      "page_no": 8, "max_score": 6, "bbox": [10, 20, 100, 40]}],
    }
    prediction = {"analysis": {
        "document": {"sha256": digest},
        "project_fields": [{"field": "project_title", "raw_value": "示例工程"},
                           {"field": "project_number", "raw_value": None}],
        "criteria": [
            {"criterion_id": "hit", "category": "technical", "source_text": "类似工程业绩每项得2分最高6分。",
             "max_score": 6, "evidence": [{"page_no": 8, "quote": "类似工程业绩每项得2分最高6分。",
                                            "bbox": [10, 20, 100, 40]}]},
            {"criterion_id": "extra", "category": "commercial", "source_text": "企业信用评价加3分",
             "max_score": 3, "evidence": [{"page_no": 12, "quote": "企业信用评价加3分"}]},
        ],
    }}

    report = evaluate(gold, prediction)

    assert report["metrics"]["project_field_exact_accuracy"] == 1.0
    assert report["metrics"]["criterion_precision"] == 0.5
    assert report["metrics"]["criterion_recall"] == 1.0
    assert report["metrics"]["criterion_max_score_accuracy"] == 1.0
    assert report["metrics"]["evidence_quote_and_page_accuracy"] == 1.0
    assert report["metrics"]["bbox_iou_at_least_0_5_accuracy"] == 1.0
    corpus = aggregate_corpus([{"project_id": "p-1", "split": "blind_test", "format": "scanned_pdf",
                               "report": report}])
    assert corpus["by_split"]["blind_test"]["metrics"]["criterion_f1"] == report["metrics"]["criterion_f1"]
    assert corpus["by_format"]["scanned_pdf"]["samples"] == 1


def test_corpus_evaluation_rejects_project_leakage_across_splits():
    sample = {"project_id": "same-project", "split": "development", "report": {"counts": {}}}
    duplicate = {**sample, "split": "blind_test"}

    try:
        aggregate_corpus([sample, duplicate])
    except ValueError as exc:
        assert "项目级泄漏" in str(exc)
    else:
        raise AssertionError("the corpus evaluator must reject project leakage between splits")


def test_gold_evaluation_rejects_wrong_document_version():
    gold = {"schema_version": "1.0", "annotation_complete": True,
            "document": {"sha256": "a" * 64},
            "coverage_reviewed": ["project_fields", "technical_criteria", "commercial_criteria", "price_criteria"],
            "project_fields": {"project_title": None}, "criteria": []}

    try:
        evaluate(gold, {"document": {"sha256": "b" * 64}, "project_fields": [], "criteria": []})
    except ValueError as exc:
        assert "SHA-256" in str(exc)
    else:
        raise AssertionError("evaluation should reject predictions from a different source file version")


def test_scoring_candidate_keeps_source_and_evidence_without_assumed_total():
    criteria = extract_criteria(_blocks(), "doc-1")

    assert len(criteria) == 1
    candidate = criteria[0]
    assert candidate.category == "technical"
    assert candidate.subcategory == "similar_projects"
    assert candidate.max_score == 6
    assert candidate.source_text == _blocks()[3]["text"]
    assert candidate.evidence[0].page_no == 4
    assert candidate.evidence[0].bbox == [30, 40, 400, 80]
    assert "不予受理" not in candidate.source_text


def test_scoring_extraction_ignores_boilerplate_mentions_outside_scoring_tables():
    blocks = [
        {"block_id": "b-1", "kind": "paragraph", "text": "评标委员会按评标办法评分", "source_index": 0},
        {"block_id": "b-2", "kind": "paragraph", "text": "类似工程业绩每项得 2 分，最高 6 分", "source_index": 1},
    ]

    assert extract_criteria(blocks, "doc-boilerplate") == []


def test_scoring_caps_handle_negative_ranges_and_composite_staffing_rows():
    blocks = [
        {"block_id": "h", "kind": "paragraph", "text": "评标办法前附表之三：商务评分标准", "source_index": 0},
        {"block_id": "h1", "kind": "table_row", "text": "序号 | 项目 | 评审内容及分值 | 项目内容",
         "cells": ["序号", "项目", "评审内容及分值", "项目内容"], "source_index": 1, "table_id": "t-1", "row_index": 0},
        {"block_id": "negative", "kind": "table_row", "text": "1 | 诚信评价 | 不良行为（-30-0） | 有不良行为扣30分",
         "cells": ["1", "诚信评价", "不良行为（-30-0）", "有不良行为扣30分"], "source_index": 2, "table_id": "t-1", "row_index": 1},
        {"block_id": "h2", "kind": "paragraph", "text": "评标办法前附表之四：技术评分标准", "source_index": 3},
        {"block_id": "h3", "kind": "table_row", "text": "序号 | 项目 | 评审内容及分值 | 项目内容",
         "cells": ["序号", "项目", "评审内容及分值", "项目内容"], "source_index": 4, "table_id": "t-2", "row_index": 0},
        {"block_id": "staff", "kind": "table_row",
         "text": "1 | 团队配置（15分） | 工作组织及人员配备（15分） | （1）本项最高得4分（2）本项最高得8分（3）本项最高得3分（4）本项最高得3分",
         "cells": ["1", "团队配置（15分）", "工作组织及人员配备（15分）",
                   "（1）本项最高得4分（2）本项最高得8分（3）本项最高得3分（4）本项最高得3分",
                   "（1）本项最高得4分（2）本项最高得8分（3）本项最高得3分（4）本项最高得3分"],
         "source_index": 5, "table_id": "t-2", "row_index": 1},
    ]

    criteria = extract_criteria(blocks, "doc-score")
    negative = next(item for item in criteria if item.criterion_label.startswith("不良行为"))
    staffing = next(item for item in criteria if item.subcategory == "team_staffing")

    assert negative.max_score == 0
    assert staffing.criterion_label == "工作组织及人员配备（15分）"
    assert staffing.max_score == 15
    assert len(staffing.conditions) == 4
    assert "合计18分" in staffing.review_reason


def test_text_upload_parser_records_hash_and_anchor(tmp_path):
    path = tmp_path / "tender.txt"
    path.write_text("项目名称：示例工程\n评标办法\n商务评分 最高 3 分", encoding="utf-8")

    parsed = parse_document(path, max_pages=100)

    assert len(parsed["sha256"]) == 64
    assert parsed["parser"] == "utf-8-text"
    assert parsed["blocks"][0]["text"] == "项目名称：示例工程"
    assert all(item["state"] == "done" for item in parsed["ledger"])


def test_xlsx_parser_preserves_sheet_cell_and_merged_context(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "保证金清单"
    sheet.merge_cells("A1:A2")
    sheet["A1"] = "包件"
    sheet["B1"] = "保证金金额"
    sheet["B2"] = "581万元"
    path = tmp_path / "guarantees.xlsx"
    workbook.save(path)

    parsed = parse_document(path, max_pages=100)
    rows = parsed["blocks"]

    assert parsed["parser"] == "openpyxl"
    assert rows[0]["sheet_name"] == "保证金清单"
    assert rows[0]["cell_range"] == "A1:B1"
    assert rows[1]["cells"] == ["包件", "581万元"]
    assert rows[1]["cell_range"] == "A2:B2"
    assert rows[1]["section_path"] == ["保证金清单"]
    assert len(parsed["sections"]) == 1
    assert parsed["ledger"][1]["sheet_name"] == "保证金清单"
    assert parsed["ledger"][1]["cell_range"] == "A2:B2"


def test_zip_bundle_parsing_keeps_nested_document_sources_and_skips_old_formats(tmp_path):
    nested = BytesIO()
    with zipfile.ZipFile(nested, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("nested.docx", _docx_bytes())
    package_doc = Document()
    package_doc.add_heading("结算审核包2", level=1)
    package_bytes = BytesIO()
    package_doc.save(package_bytes)
    workbook = Workbook()
    workbook.active["A1"] = "保证金清单"
    spreadsheet_bytes = BytesIO()
    workbook.save(spreadsheet_bytes)
    path = tmp_path / "bundle.zip"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("main.docx", _docx_bytes())
        archive.writestr("nested.zip", nested.getvalue())
        archive.writestr("技术规范书.docx", package_bytes.getvalue())
        archive.writestr("保证金清单.xlsx", spreadsheet_bytes.getvalue())
        archive.writestr("legacy.doc", b"legacy Word")

    parsed = parse_document(path, max_pages=100, ocr_enabled=False)

    assert parsed["archive_stats"]["documents"] == 4
    assert parsed["scope_hint"] == "结算审核包2"
    assert len({block["block_id"] for block in parsed["blocks"]}) == len(parsed["blocks"])
    assert {block["source_file"].split("!/")[-1] for block in parsed["blocks"]} == {
        "main.docx", "nested.docx", "技术规范书.docx", "保证金清单.xlsx"
    }
    assert any(item["kind"] == "doc" and item["state"] == "failed_review" for item in parsed["ledger"])


def test_package_specific_fields_match_package_row_and_keep_table_evidence():
    blocks = [
        {"block_id": "b-head", "kind": "table_row", "text": "包号 | 工程名称 | 包名称 | 项目地点 | 招标范围 | 开始时间（年月） | 完成时间（年月） | 分项限价（万元） | 最高限价（万元）",
         "cells": ["包号", "工程名称", "包名称", "项目地点", "招标范围", "开始时间（年月）", "完成时间（年月）", "分项限价（万元）", "最高限价（万元）"],
         "source_index": 0, "table_id": "t-1", "row_index": 0},
        {"block_id": "b-package1", "kind": "table_row", "text": "包1 | 浙江特高压交流环网线路工程 | 浙江环网结算审核包1 | 浙江 | 审核 | 2026年11月 | 2030年4月 | \\ | 581",
         "cells": ["包1", "浙江特高压交流环网线路工程", "浙江环网结算审核包1", "浙江", "审核", "2026年11月", "2030年4月", "\\", "581"],
         "source_index": 1, "table_id": "t-1", "row_index": 1, "source_file": "notice.docx"},
        {"block_id": "b-package12", "kind": "table_row", "text": "包12 | 其他线路工程 | 其他结算审核包12 | 浙江 | 审核 | 2026年11月 | 2030年4月 | \\ | 700",
         "cells": ["包12", "其他线路工程", "其他结算审核包12", "浙江", "审核", "2026年11月", "2030年4月", "\\", "700"],
         "source_index": 2, "table_id": "t-1", "row_index": 2, "source_file": "notice.docx"},
    ]

    fields = {item.field: item for item in _package_specific_fields(blocks, "doc-1", "结算审核包1")}

    assert fields["project_title"].raw_value == "浙江特高压交流环网线路工程"
    assert fields["package_name"].raw_value == "浙江环网结算审核包1"
    assert fields["price_limit"].normalized_value == "5810000.00"
    assert fields["duration"].raw_value == "开始时间 2026年11月 至 完成时间 2030年4月"
    assert fields["price_limit"].evidence[0].block_id == "b-package1"


def test_upload_endpoint_persists_task_and_review_uses_history(tmp_path, monkeypatch):
    test_engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(test_engine)
    TestSession = sessionmaker(bind=test_engine, autoflush=False, expire_on_commit=False)

    def override_session():
        session = TestSession()
        try:
            yield session
        finally:
            session.close()

    monkeypatch.setattr(settings, "data_dir", tmp_path)
    monkeypatch.setattr("bidreader.app.submit", lambda _run_id: None)
    app.dependency_overrides[get_session] = override_session
    try:
        with TestClient(app) as client:
            response = client.post("/api/v1/tenders", files={"file": ("tender.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
            assert response.status_code == 202
            ids = response.json()
            with TestSession() as session:
                run = session.get(Run, ids["run_id"])
                assert run is not None and run.sha256 == ids["sha256"]
                assert session.scalar(select(RunEvent).where(RunEvent.run_id == run.id)) is not None
                run.status = "failed"
                session.commit()
            retried = client.post(f"/api/v1/runs/{ids['run_id']}/retry")
            assert retried.status_code == 202
            assert retried.json()["status"] == "queued"
            with TestSession() as session:
                run = session.get(Run, ids["run_id"])
                run.status = "needs_review"
                payload = {"run_id": run.id, "document": {"document_id": "d-1", "blocks": [
                    {"block_id": "b-1", "text": "原文评分条款", "page_no": 1, "bbox": None}]}, "project_fields": [], "sections": [], "criteria": [
                    {"criterion_id": "c-1", "category": "technical", "subcategory": None, "source_text": "原文评分条款", "score_text": None,
                     "max_score": 3, "tiers": [], "evidence_required": [], "conditions": [], "formula_text": None,
                     "evidence": [{"document_id": "d-1", "page_no": 1, "block_id": "b-1", "quote": "原文评分条款", "bbox": None, "source_kind": "text"}],
                     "extraction_state": "candidate", "review_state": "pending", "review_reason": None}], "ledger": [], "warnings": [], "parser_versions": {}}
                session.add(Analysis(run_id=run.id, payload=payload))
                session.commit()
            reviewed = client.patch(f"/api/v1/runs/{ids['run_id']}/criteria/c-1", json={"action": "confirmed", "reviewer": "tester", "note": "source checked"})
            assert reviewed.status_code == 200
            history = client.get(f"/api/v1/runs/{ids['run_id']}/reviews")
            assert history.status_code == 200
            assert history.json()[0]["action"] == "confirmed"
            assert history.json()[0]["before"]["review_state"] == "pending"
    finally:
        app.dependency_overrides.clear()
        Base.metadata.drop_all(test_engine)
        test_engine.dispose()


def test_similar_project_rows_keep_parent_and_specific_work_type():
    blocks = [
        {"block_id": "b-heading", "kind": "paragraph", "text": "技术评分标准"},
        {"block_id": "b-header", "kind": "table_row", "text": "序号 | 项目 | 评审内容及分值 | 项目内容 | 分值",
         "cells": ["序号", "项目", "评审内容及分值", "项目内容", "分值"], "table_id": "t-1", "row_index": 0},
        {"block_id": "b-1", "kind": "table_row", "text": "1 | 技术水平 | 相似工程业绩（30分） | 800kV及以上工程结算审核 | 每项得4分",
         "cells": ["1", "技术水平", "相似工程业绩（30分）", "800kV及以上工程结算审核", "每项得4分"],
         "table_id": "t-1", "row_index": 1},
        {"block_id": "b-2", "kind": "table_row", "text": "2 | 技术水平 | 相似工程业绩（30分） | 500kV工程结算审核 | 每项得2分",
         "cells": ["2", "技术水平", "相似工程业绩（30分）", "500kV工程结算审核", "每项得2分"],
         "table_id": "t-1", "row_index": 2},
    ]

    criteria = extract_criteria(blocks, "doc-1")

    assert [(item.parent_label, item.criterion_label, item.max_score) for item in criteria] == [
        ("相似工程业绩（30分）", "800kV及以上工程结算审核", 4.0),
        ("相似工程业绩（30分）", "500kV工程结算审核", 2.0),
    ]
